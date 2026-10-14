"""Consulte la base de données.

    python -m bodacc.report vehicles                            # les 30 derniers véhicules vus
    python -m bodacc.report vehicles --max-price 20000 --max-km 50000
    python -m bodacc.report vehicles --marque renault --min-year 2018 --sort km
    python -m bodacc.report vehicles --cabinet lga --carburant diesel --lieu 33220
    python -m bodacc.report vehicles --marque renault --since 2026-09-19 --limit 100
    python -m bodacc.report vehicles --csv data/exports/vehicules.csv # tout exporter (ouvrable dans Excel)
    python -m bodacc.report vehicles --destination "87000 Limoges"   # + distance et coût de transport estimé (fourchette)
    python -m bodacc.report unmatched                           # liquidateurs sans cabinet du registre

La colonne « cote » vient des cotes saisies à la main avec `python -m bodacc.cote`.
"""
import argparse
import csv
from collections import Counter
from pathlib import Path

from bodacc.actif_filter import normalize
from bodacc.cabinets import cabinet_name, match_cabinet, normalize_cabinet
from bodacc.gemweb_rss import CabinetRegistry
from bodacc.paths import DB_PATH
from bodacc.storage import Storage
from bodacc.update import is_vehicle
from bodacc.vehicle_filters import filter_vehicles, sort_vehicles
from bodacc.vehicles import count_vehicles

SORTS = ["recent", "price", "estimation", "km", "year", "cote"]
TABLE_COLUMNS = [
    ("first_seen", "vu le"), ("cabinet", "cabinet"), ("nb_vehicules", "nb"), ("marque", "marque"), ("annee", "année"),
    ("immatriculation", "plaque"), ("carburant", "carburant"),
    ("kilometrage", "km"), ("estimation", "estimation"), ("cote", "cote €"), ("lieu", "lieu"),
]
# Dans les exports, « estimation » est le montant en nombre (triable dans Excel) ; le texte du cabinet n'est gardé que
# quand il n'y a pas de montant (« Non renseigné »).
CSV_COLUMNS = [key for key, _ in TABLE_COLUMNS] + [
    "date_limite_offre", "cote_source", "titre", "description", "url_detail"]


def sort_option(text):
    """« km,estimation,cote » : une ou plusieurs clés de SORTS, séparées par des virgules."""
    keys = [k.strip() for k in text.split(",") if k.strip()]
    unknown = [k for k in keys if k not in SORTS]
    if not keys or unknown:
        raise argparse.ArgumentTypeError(f"tri inconnu : {', '.join(unknown) or text!r} (choix : {', '.join(SORTS)})")
    return ",".join(keys)


def for_export(row):
    """La ligne avec l'estimation en nombre quand le texte du cabinet contient un montant."""
    price = row.get("prix")
    if price is None:
        return row
    return {**row, "estimation": int(price) if float(price).is_integer() else price}
LEADS_CSV_COLUMNS = [
    "cabinet_nom", "nb_liquidations_cabinet", "entreprise", "siren", "ville", "cp", "tribunal",
    "date_jugement", "nature", "liquidateur", "url_bodacc", "first_seen",
]


def vehicle_rows(storage, since=None, marque=None):
    """Véhicules en base, les plus récemment vus d'abord.

    since : date AAAA-MM-JJ (les dates ISO se comparent comme du texte).
    marque: filtre insensible à la casse et aux accents ("citroen" trouve Citroën).
    """
    rows = [dict(r) for r in storage.new_since("actifs", since or "0") if is_vehicle(r)]
    cotes, fiches = storage.cotes(), storage.fiches()
    for row in rows:
        if row.get("nb_vehicules") is None:     # lignes stockées avant l'ajout de la colonne : on la calcule
            row["nb_vehicules"] = count_vehicles(row.get("titre"), row.get("description"))
        if not row.get("date_limite_offre"):    # absente du flux : celle lue sur la fiche (bodacc.verify), s'il y en a une
            row["date_limite_offre"] = fiches.get(row["guid"], {}).get("date_limite")
        cote = cotes.get(row["guid"], {})
        row["cote"], row["cote_source"] = cote.get("valeur"), cote.get("source")
    if marque:
        wanted = normalize(marque)
        rows = [r for r in rows if wanted in normalize(r.get("marque"))]
    return rows


def unmatched_liquidations(storage, registry):
    """Liquidations dont le liquidateur n'est pas au registre : autant de pistes à contacter.

    Chaque ligne reçoit 'cabinet_nom' (le nom lisible, tel qu'écrit au BODACC) et
    'nb_liquidations_cabinet'. Tri : cabinets les plus fréquents d'abord.

    On refait la correspondance avec le registre actuel : ajouter un cabinet à
    cabinets.json le fait disparaître de cette liste sans relancer le BODACC.
    """
    names = registry.names()
    keyed = [
        (normalize_cabinet(cabinet_name(row["liquidateur"])), row)
        for row in storage.new_since("liquidations", "0")
        if row["liquidateur"] and not match_cabinet(row["liquidateur"], names)
    ]
    counts = Counter(key for key, _ in keyed)
    rows = [
        {**row, "cabinet_nom": cabinet_name(row["liquidateur"]),
         "nb_liquidations_cabinet": counts[key], "_cle": key}
        for key, row in keyed
    ]
    rows.sort(key=lambda r: (-r["nb_liquidations_cabinet"], r["_cle"], r["entreprise"] or ""))
    return rows


def unmatched_cabinets(storage, registry):
    """[(nom de cabinet, nombre de liquidations)] pour les liquidateurs absents du registre."""
    counts = Counter(r["_cle"] for r in unmatched_liquidations(storage, registry))
    return counts.most_common()


def format_table(rows, columns=TABLE_COLUMNS, width=30, link_key=None):
    """Tableau texte aligné, sans dépendance externe.

    link_key : clé d'une URL (ex : 'url_detail'). Elle est écrite en entier sous chaque ligne,
    et non en colonne : une URL est trop longue pour un tableau et serait coupée.
    """
    def cell(key, value):
        if value is None:
            return ""
        if key == "first_seen":
            return str(value)[:10]
        if key in ("kilometrage", "cote"):
            return f"{value:,}".replace(",", " ")
        return str(value)

    lines = [[header for _, header in columns]]
    lines += [[cell(key, row.get(key)) for key, _ in columns] for row in rows]
    lines = [[c[:width] for c in line] for line in lines]
    widths = [max(len(line[i]) for line in lines) for i in range(len(columns))]

    out = ["  ".join(c.ljust(w) for c, w in zip(line, widths)).rstrip() for line in lines]
    out.insert(1, "  ".join("-" * w for w in widths))

    if link_key:
        with_links = out[:2]                                  # en-tête et séparateur
        for row, line in zip(rows, out[2:]):
            with_links.append(line)
            if row.get(link_key):
                with_links.append(f"    {row[link_key]}")
        out = with_links
    return "\n".join(out)


def write_csv(rows, path, columns=CSV_COLUMNS):
    """CSV en UTF-8 avec BOM : Excel affiche correctement les accents."""
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return path


TRANSPORT_TABLE = [("km_route", "km route"), ("transport", "transport €")]
TRANSPORT_CSV = ["km_route", "transport_min", "transport_max"]
MARKET_TABLE = [("marche", "marché € (n)")]
MARKET_CSV = ["marche_median", "marche_bas", "marche_haut", "marche_n", "marche_critere"]
XLSX_HEADERS = {**dict(TABLE_COLUMNS), "estimation": "estimation €", "date_limite_offre": "date limite offre", "url_detail": "lien",
                "km_route": "km route", "transport": "transport €", "transport_min": "transport min €", "transport_max": "transport max €",
                "marche_median": "marché médiane €", "marche_bas": "marché bas €", "marche_haut": "marché haut €", "marche_n": "annonces comparées",
                "marche_critere": "critère de comparaison", "cote_source": "source de la cote"}
XLSX_WIDTHS = {"cabinet": 22, "lieu": 30, "titre": 45, "description": 70, "estimation": 16, "url_detail": 40}


def write_xlsx(rows, path, columns=CSV_COLUMNS):
    """Classeur Excel : en-têtes en gras et figés, filtres, colonnes ajustées, liens cliquables, km et prix numériques."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Véhicules"
    sheet.append([XLSX_HEADERS.get(key, key) for key in columns])
    for row in rows:
        sheet.append([str(row.get(key))[:10] if key == "first_seen" and row.get(key) else row.get(key) for key in columns])
    for number, key in enumerate(columns, start=1):
        letter = get_column_letter(number)
        sheet.column_dimensions[letter].width = XLSX_WIDTHS.get(key, 12)
        for cell in sheet[letter][1:]:
            if key == "url_detail" and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
            elif key in ("kilometrage",):
                cell.number_format = "#,##0"
            elif key in ("estimation", "cote") and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0 "€"'
            elif key == "description":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def write_xlsx_or_explain(rows, path, columns=CSV_COLUMNS):
    try:
        return write_xlsx(rows, path, columns)
    except PermissionError:
        raise SystemExit(f"Impossible d'écrire {path} : le fichier est probablement ouvert dans Excel.\n"
                         "Fermez-le puis relancez, ou choisissez un autre nom avec --xlsx.")


def write_csv_or_explain(rows, path, columns=CSV_COLUMNS):
    """Comme write_csv, mais un fichier verrouillé donne un message clair, pas une trace."""
    try:
        return write_csv(rows, path, columns)
    except PermissionError:
        raise SystemExit(
            f"Impossible d'écrire {path} : le fichier est probablement ouvert dans Excel.\n"
            "Fermez-le puis relancez, ou choisissez un autre nom avec --csv."
        )


def main(argv=None, geocoder=None, market=None):
    parser = argparse.ArgumentParser(description="Consulte la base BODACC.")
    parser.add_argument("--db", default=DB_PATH, help="fichier SQLite (défaut : data/bodacc.db)")
    commands = parser.add_subparsers(dest="command", required=True)

    veh = commands.add_parser("vehicles", help="véhicules en vente")
    veh.add_argument("--since", help="vus pour la première fois à partir de AAAA-MM-JJ")
    veh.add_argument("--marque", help="filtrer sur une marque (ex : renault)")
    veh.add_argument("--limit", type=int, default=30, help="lignes affichées (défaut : 30)")
    veh.add_argument("--csv", help="exporter tous les résultats dans ce fichier CSV")
    veh.add_argument("--xlsx", help="exporter tous les résultats dans ce classeur Excel (.xlsx)")
    veh.add_argument("--cabinet", help="cabinet contenant ce texte (ex : lga)")
    veh.add_argument("--carburant", help="carburant (ex : diesel, essence, electrique, hybride)")
    veh.add_argument("--lieu", help="lieu contenant ce texte : ville ou code postal (ex : 33220, paris)")
    veh.add_argument("--min-year", type=int, help="année minimum")
    veh.add_argument("--max-year", type=int, help="année maximum")
    veh.add_argument("--min-price", type=float, help="prix minimum en euros")
    veh.add_argument("--max-price", type=float, help="prix maximum en euros")
    veh.add_argument("--min-km", type=int, help="kilométrage minimum")
    veh.add_argument("--max-km", type=int, help="kilométrage maximum")
    veh.add_argument("--destination", help="lieu de livraison (ex : « 87000 Limoges ») : ajoute la distance et une fourchette de coût de transport")
    veh.add_argument("--market", action="store_true",
                     help="ajoute le prix de marché (médiane d'annonces Leboncoin comparables, via Carapis ; clé dans .env)")
    veh.add_argument("--no-links", action="store_true",
                     help="ne pas afficher le lien de l'annonce sous chaque véhicule")
    veh.add_argument("--include-unknown", action="store_true",
                     help="garder les véhicules dont prix, km, année ou carburant est inconnu")
    veh.add_argument("--sort", type=sort_option, default="recent",
                     help="recent (défaut), estimation (= price), km, cote : croissant ; year : récent d'abord. "
                          "Plusieurs clés séparées par des virgules : km,estimation,cote")

    unm = commands.add_parser("unmatched", help="liquidateurs absents du registre")
    unm.add_argument("--limit", type=int, default=25)
    unm.add_argument("--csv", help="exporter les liquidations (pistes) dans ce fichier CSV")

    args = parser.parse_args(argv)

    with Storage(args.db) as storage:
        if args.command == "vehicles":
            criteria = dict(
                cabinet=args.cabinet, carburant=args.carburant, lieu=args.lieu,
                min_year=args.min_year, max_year=args.max_year,
                min_price=args.min_price, max_price=args.max_price,
                min_km=args.min_km, max_km=args.max_km)
            found = vehicle_rows(storage, args.since, args.marque)
            rows = sort_vehicles(
                filter_vehicles(found, **criteria, keep_unknown=args.include_unknown), args.sort)

            note = ""
            if not args.include_unknown:
                # Un filtre écarte les valeurs inconnues : on dit combien, pour que ce ne soit pas une surprise.
                hidden = len(filter_vehicles(found, **criteria, keep_unknown=True)) - len(rows)
                if hidden:
                    note = (f"\n{hidden} autre(s) véhicule(s) écarté(s) car une valeur est inconnue "
                            "(--include-unknown pour les garder)")

            table_columns, export_columns = TABLE_COLUMNS, CSV_COLUMNS
            exporting = bool(args.csv or args.xlsx)
            if args.destination:
                from bodacc import transport        # ici et pas en tête : prospects.search (que transport utilise) importe ce module
                try:
                    transport.annotate(rows if exporting else rows[: args.limit], args.destination, geocoder or transport.Geocoder())
                except transport.TransportError as error:
                    raise SystemExit(str(error))
                table_columns, export_columns = TABLE_COLUMNS + TRANSPORT_TABLE, CSV_COLUMNS + TRANSPORT_CSV
                note += (f"\nTransport vers {args.destination} : fourchette estimée pour une voiture standard qui roule (convoyage, tout compris), "
                         "pas un devis (voir bodacc/transport.py).")
            if args.market:
                from bodacc import market as market_module
                try:
                    skipped = market_module.annotate(rows if exporting else rows[: args.limit], market or market_module.Market())
                except market_module.MarketError as error:
                    raise SystemExit(f"Carapis : {error}")
                table_columns, export_columns = table_columns + MARKET_TABLE, export_columns + MARKET_CSV
                note += ("\nMarché : médiane des prix demandés sur un échantillon d'annonces Leboncoin comparables (marque, modèle, "
                         "année ±2, km ±40 %), n = nombre d'annonces ; critère affiché (« année », « modèle seul ») quand la recherche a dû être élargie ; "
                         "vide si marque/modèle inconnus, lot de plusieurs véhicules ou moins de 3 annonces. "
                         "Pas la cote La Centrale (voir bodacc/market.py).")
                if skipped:
                    note += f"\n{skipped} véhicule(s) non estimé(s) : limite de requêtes atteinte pour cette exécution, relancez plus tard."

            if exporting:
                exported = [for_export(r) for r in rows]
                if args.csv:
                    write_csv_or_explain(exported, args.csv, export_columns)
                if args.xlsx:
                    write_xlsx_or_explain(exported, args.xlsx, export_columns)
                print(f"{len(rows)} véhicule(s) exporté(s) dans {' et '.join(p for p in (args.csv, args.xlsx) if p)}{note}")
            else:
                print(format_table(rows[: args.limit], table_columns, link_key=None if args.no_links else "url_detail"))
                total = sum(r["nb_vehicules"] for r in rows)
                count = f"{len(rows)} véhicule(s)" if total == len(rows) else f"{len(rows)} annonce(s), {total} véhicules"
                print(f"\n{count}" + (f", {args.limit} affiché(s)" if len(rows) > args.limit else "") + note)
        elif args.csv:
            leads = unmatched_liquidations(storage, CabinetRegistry())
            write_csv_or_explain(leads, args.csv, LEADS_CSV_COLUMNS)
            print(f"{len(leads)} liquidation(s) hors registre exportée(s) dans {args.csv}")
        else:
            missing = unmatched_cabinets(storage, CabinetRegistry())
            print(f"{'liquidations':>12}  cabinet (à ajouter au registre ?)")
            for name, count in missing[: args.limit]:
                print(f"{count:>12}  {name}")
            print(f"\n{sum(c for _, c in missing)} liquidation(s), {len(missing)} nom(s) distinct(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
