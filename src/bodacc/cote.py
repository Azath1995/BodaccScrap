"""Saisit à la main la cote d'un véhicule (lue sur La Centrale, L'argus...) : elle apparaît dans `bodacc.report`.

    python -m bodacc.cote FX-505-JK 24408 --source "La Centrale : Master III (2) L2H2 Blue dCi 145 Confort"
    python -m bodacc.cote FX-505-JK --supprimer
    python -m bodacc.cote --liste
    python -m bodacc.cote --importer data/exports/vehicules.xlsx     # cotes tapées dans l'export Excel

Le véhicule est désigné par sa plaque (tirets et espaces ignorés) ou par l'URL de l'annonce. La cote est rangée dans
la table `cotes` de la base, que `bodacc.update` ne modifie jamais : elle survit aux mises à jour et aux exports.

`--importer` relit les colonnes « cote € » et « source de la cote » d'un export de `bodacc.report` (ligne retrouvée
par sa colonne « lien »). Une case de cote vide ne supprime rien (pour cela : --supprimer). Les autres colonnes
viennent des cabinets et seraient réécrites à la mise à jour suivante : elles ne sont pas importées.
"""
import argparse
import re

from bodacc.paths import DB_PATH
from bodacc.storage import Storage

IMPORT_HEADERS = {"lien": "url", "cote €": "valeur", "source de la cote": "source"}


def read_cotes(path):
    """[(url, valeur, source)] des lignes de l'export Excel qui ont une cote."""
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [IMPORT_HEADERS.get(str(h).strip()) if h else None for h in next(rows)]
    if "url" not in headers or "valeur" not in headers:
        raise SystemExit(f"{path} : colonnes « lien » et « cote € » introuvables (est-ce un export de bodacc.report ?)")
    found = []
    for line in rows:
        row = {key: value for key, value in zip(headers, line) if key}
        value = row.get("valeur")
        if value in (None, ""):
            continue
        try:
            value = int(round(float(str(value).replace("€", "").replace(" ", "").replace(" ", "").replace(",", "."))))
        except ValueError:
            raise SystemExit(f"Cote illisible {row.get('valeur')!r} pour {row.get('url')} : mettez un nombre (ex : 24408).")
        found.append((row.get("url"), value, row.get("source") or None))
    return found


def import_cotes(storage, path):
    """Enregistre les cotes du fichier. Retourne (nouvelles ou modifiées, inchangées, liens inconnus)."""
    known = {r[0]: r[1] for r in storage.conn.execute("SELECT url_detail, guid FROM actifs")}
    current = storage.cotes()
    changed, same, unknown = [], 0, []
    for url, value, source in read_cotes(path):
        guid = known.get(url)
        if guid is None:
            unknown.append(url)
            continue
        before = current.get(guid, {})
        source = source or before.get("source")
        if before.get("valeur") == value and before.get("source") == source:
            same += 1
            continue
        storage.set_cote(guid, value, source)
        changed.append((url, before.get("valeur"), value))
    return changed, same, unknown


def plate_key(text):
    return re.sub(r"[\s-]", "", text or "").upper()


def find_actif(storage, reference):
    """L'actif désigné par une plaque ou une URL d'annonce ; SystemExit si aucun ou plusieurs."""
    rows = [dict(r) for r in storage.conn.execute("SELECT guid, titre, immatriculation, url_detail FROM actifs")]
    if reference.startswith("http"):
        found = [r for r in rows if reference in (r["guid"], r["url_detail"])]
    else:
        found = [r for r in rows if r["immatriculation"] and plate_key(r["immatriculation"]) == plate_key(reference)]
    if not found:
        raise SystemExit(f"Aucun véhicule {reference!r} dans la base.")
    if len(found) > 1:
        listing = "\n".join(f"  {r['url_detail'] or r['guid']}  {r['titre']}" for r in found)
        raise SystemExit(f"{len(found)} annonces portent la plaque {reference} : donnez l'URL de l'une d'elles.\n{listing}")
    return found[0]


def main(argv=None):
    parser = argparse.ArgumentParser(description="Cote saisie à la main pour un véhicule de la base.")
    parser.add_argument("vehicule", nargs="?", help="plaque (ex : FX-505-JK) ou URL de l'annonce")
    parser.add_argument("valeur", nargs="?", type=int, help="cote en euros (ex : 24408)")
    parser.add_argument("--source", help="d'où vient la cote (site, version choisie)")
    parser.add_argument("--supprimer", action="store_true", help="retire la cote de ce véhicule")
    parser.add_argument("--liste", action="store_true", help="affiche toutes les cotes saisies")
    parser.add_argument("--importer", metavar="FICHIER.xlsx", help="enregistre les cotes tapées dans un export Excel")
    parser.add_argument("--db", default=DB_PATH)
    args = parser.parse_args(argv)

    with Storage(args.db) as storage:
        if args.importer:
            try:
                changed, same, unknown = import_cotes(storage, args.importer)
            except PermissionError:
                raise SystemExit(f"Impossible de lire {args.importer} : fermez-le dans Excel puis relancez.")
            titles = {r[0]: r[1] for r in storage.conn.execute("SELECT url_detail, titre FROM actifs")}
            print(f"{len(changed)} cote(s) enregistrée(s), {same} déjà à jour")
            for url, before, value in changed:
                old = f"{before:,} € -> ".replace(",", " ") if before is not None else ""
                new = f"{value:,} €".replace(",", " ")
                print(f"   {old}{new}  {(titles.get(url) or '').removeprefix('A vendre : ')}")
            if unknown:
                print(f"{len(unknown)} ligne(s) dont l'annonce n'est plus en base (ignorées) :")
                for url in unknown:
                    print(f"   {url}")
            return 0
        if args.liste:
            titles = {r[0]: (r[1], r[2]) for r in storage.conn.execute("SELECT guid, titre, immatriculation FROM actifs")}
            for guid, cote in storage.cotes().items():
                titre, plaque = titles.get(guid, ("(annonce absente)", None))
                value = f"{cote['valeur']:>8,} €".replace(",", " ")
                print(f"{value}  {plaque or '':10}  {cote['date']}  {titre}  {cote['source'] or ''}")
            return 0
        if not args.vehicule or (args.valeur is None and not args.supprimer):
            parser.error("donnez un véhicule et une valeur (ou --supprimer, ou --liste)")
        actif = find_actif(storage, args.vehicule)
        storage.set_cote(actif["guid"], None if args.supprimer else args.valeur, args.source)
        done = "supprimée" if args.supprimer else f"{args.valeur:,} €".replace(",", " ")
        print(f"Cote {done} : {actif['titre']} ({actif['immatriculation'] or actif['guid']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
