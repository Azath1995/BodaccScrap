"""Consulte la base espagnole. Mêmes commandes et mêmes options que côté France (`python -m bodacc.report`).

    python -m es.report vehicles                              # véhicules à vendre (ouverts ou à venir)
    python -m es.report vehicles --max-price 20000 --marque mitsubishi --sort price
    python -m es.report vehicles --lieu madrid --status all --csv data/es/exports/vehicules.csv
    python -m es.report unmatched                             # administrateurs concursaux (aucun registre espagnol : tous)
    python -m es.report liquidations --type liquidation       # ouvertures de liquidation du BORME
    python -m es.report liquidations --lieu madrid --since 2026-09-01 --csv data/es/exports/liquidations.csv
    python -m es.report links                                 # lots rattachés à un acte du BORME (donc à un administrateur)
"""
import argparse
import csv
import re
import unicodedata
from collections import Counter
from pathlib import Path

from es.paths import DB_PATH
from es.storage import COLUMNS, LOTE_COLUMNS, Storage

TYPES = {
    "declaration": "Auto de declaración de concurso%",
    "liquidation": None,                       # se filtre sur la colonne liquidacion
    "conclusion": "Auto de conclusión%",
}
STATUSES = {"live": ("abierta", "proxima"), "open": ("abierta",), "upcoming": ("proxima",), "closed": ("cerrada",), "all": None}
SORTS = ["recent", "price"]
LIQUIDATION_TABLE = [("fecha_publicacion", "publié"), ("provincia", "province"), ("empresa", "entreprise"), ("procedimiento", "procédure"),
                     ("tipo", "acte"), ("administradores", "administrateur")]
VEHICLE_TABLE = [("vu_le", "vu le"), ("etat", "état"), ("nb", "nb"), ("marque", "marque"), ("matricula", "plaque"), ("titulo", "lot"),
                 ("valoracion", "valorisation €"), ("lieu", "lieu"), ("procedimiento", "procédure")]
LINK_TABLE = [("titulo", "lot"), ("estado", "état"), ("valoracion", "valor. €"), ("procedimiento", "procédure"), ("candidatos", "cand."),
              ("empresa", "entreprise (BORME)"), ("administradores", "administrateur (BORME)")]
STATE_LABELS = {"abierta": "ouvert", "proxima": "à venir", "cerrada": "clos"}
MAKES = ["mercedes benz", "mercedes", "land rover", "alfa romeo", "volkswagen", "renault", "peugeot", "citroen", "citroën", "seat",
         "ford", "opel", "fiat", "toyota", "nissan", "mitsubishi", "audi", "bmw", "dacia", "kia", "hyundai", "honda", "mazda",
         "skoda", "volvo", "suzuki", "iveco", "man", "scania", "daf", "piaggio", "yamaha", "harley", "kawasaki", "mini", "jeep",
         "smart", "chevrolet", "porsche", "lexus", "subaru", "tesla", "cupra", "ssangyong", "isuzu", "jcb", "caterpillar", "komatsu"]


def fold(text):
    decomposed = unicodedata.normalize("NFKD", text or "")
    return "".join(c for c in decomposed if not unicodedata.combining(c)).lower()


def guess_make(title):
    """Marque citée dans le titre d'un lot (« Vehículo Renault Master en Madrid » -> « Renault »), ou None."""
    folded = fold(title)
    for make in MAKES:
        if re.search(rf"(?<![a-z]){re.escape(fold(make))}(?![a-z])", folded):
            return make.title() if make != "mercedes benz" else "Mercedes Benz"
    return None


NUMBER_WORDS = {"un": 1, "una": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5, "seis": 6, "siete": 7, "ocho": 8,
                "nueve": 9, "diez": 10, "once": 11, "doce": 12}
VEHICLE_NOUNS = (r"(?:vehiculos?|coches?|turismos?|furgonetas?|camion(?:es)?|camionetas?|motos?|motocicletas?|"
                 r"ciclomotor(?:es)?|tractor(?:es)?|remolques?|autobus(?:es)?|autocar(?:es)?|todoterrenos?)")
COUNTED = re.compile(rf"(?<![a-z0-9])(\d{{1,3}}|{'|'.join(NUMBER_WORDS)})\s+{VEHICLE_NOUNS}(?![a-z])")


def count_vehicles(title):
    """Véhicules d'un lot : 1 sauf si le titre en annonce plusieurs (« Lote de cinco vehículos », « 3 furgonetas »).

    « Lote 1 » est un numéro de lot, pas un nombre de véhicules : seul un nombre devant un nom de véhicule compte.
    « Vehículos » sans nombre reste à 1 : on n'invente pas de nombre.
    """
    numbers = [int(NUMBER_WORDS.get(word) or word) for word in COUNTED.findall(fold(title))]
    return max([1, *numbers])


def normalize_procedure(text):
    """'000360/2022' -> '360/2022' ; '1192.5/2025' -> '1192/2025' : zéros initiaux et sous-numéro ne comptent pas."""
    number, _, year = (text or "").partition("/")
    number = number.split(".")[0]
    return f"{number.lstrip('0') or '0'}/{year}" if year else (text or "")


# ----------------------------------------------------------------------------------------------------------------------
# Liquidations (BORME)
# ----------------------------------------------------------------------------------------------------------------------

def select_liquidations(storage, type=None, lieu=None, since=None, entreprise=None):
    where, params = ["1=1"], []
    if type == "liquidation":
        where.append("liquidacion = 1")
    elif type:
        where.append("tipo LIKE ?")
        params.append(TYPES[type])
    if since:
        where.append("fecha_publicacion >= ?")
        params.append(since)
    rows = storage.rows(" AND ".join(where), params)
    if lieu:
        rows = [r for r in rows if fold(lieu) in fold(r["provincia"])]
    if entreprise:
        rows = [r for r in rows if fold(entreprise) in fold(r["empresa"])]
    return rows


def top_administrators(storage, since=None):
    """[(administrateur, nombre de procédures distinctes)] : le « qui liquide quoi » espagnol."""
    seen = {}
    for row in select_liquidations(storage, since=since):
        for admin in split_administrators(row["administradores"]):
            name, nif = split_admin(admin)
            entry = seen.setdefault(nif or fold(name), {"names": Counter(), "nif": nif, "procs": set()})
            entry["names"][name] += 1
            entry["procs"].add((row["empresa_num"], row["provincia"], row["procedimiento"]))
    ranking = []
    for entry in seen.values():
        top = max(entry["names"].values())
        name = min(n for n, c in entry["names"].items() if c == top)          # l'écriture la plus fréquente ; à égalité, la première par ordre alphabétique
        ranking.append((f"{name} ({entry['nif']})" if entry["nif"] else name, len(entry["procs"])))
    return sorted(ranking, key=lambda kv: (-kv[1], kv[0]))


def split_administrators(text):
    return [a for a in (text or "").split("; ") if a]


def split_admin(admin):
    """'PKF ATTEST CONCURSAL SLP (B95669628)' -> ('PKF ATTEST CONCURSAL SLP', 'B95669628') ; NIF None s'il n'y en a pas."""
    found = re.match(r"^(.*?)\s*\(([0-9A-Z]{9})\)$", admin)
    return (found.group(1), found.group(2)) if found else (admin, None)


# ----------------------------------------------------------------------------------------------------------------------
# Véhicules (lots eactivos)
# ----------------------------------------------------------------------------------------------------------------------

def vehicle_rows(storage, status="live", since=None, marque=None, lieu=None):
    """Lots de véhicules, les plus récemment vus d'abord ; ajoute les clés d'affichage marque, lieu, vu_le, etat."""
    rows = []
    for lot in storage.rows_lotes(order="first_seen DESC, codigo"):
        if STATUSES[status] is not None and lot["estado"] not in STATUSES[status]:
            continue
        if since and lot["first_seen"][:10] < since:
            continue
        if marque and fold(marque) not in fold(lot["titulo"]):
            continue
        place = lot["ubicacion"] or ""
        if lieu and fold(lieu) not in fold(f"{place} {lot['juzgado'] or ''} {lot['titulo'] or ''}"):
            continue
        rows.append({**lot, "marque": guess_make(lot["titulo"]), "nb": count_vehicles(lot["titulo"]), "lieu": place, "vu_le": lot["first_seen"][:10],
                     "etat": STATE_LABELS.get(lot["estado"], lot["estado"]), "url_detail": lot["url"]})
    return rows


def filter_price(rows, min_price=None, max_price=None, keep_unknown=False):
    """Bornes inclusives sur la valorisation ; une valeur inconnue est écartée, sauf keep_unknown (comme en France)."""
    if min_price is None and max_price is None:
        return rows
    kept = []
    for row in rows:
        price = row["valoracion"]
        if price is None:
            if keep_unknown:
                kept.append(row)
            continue
        if (min_price is None or price >= min_price) and (max_price is None or price <= max_price):
            kept.append(row)
    return kept


def sort_vehicles(rows, sort="recent"):
    if sort == "price":
        return sorted(rows, key=lambda r: (r["valoracion"] is None, r["valoracion"] or 0))
    return rows


COURT_STOPWORDS = {"num", "numero", "n", "juzgado", "de", "del", "la", "lo", "el", "en", "mercantil", "seccion", "tribunal",
                   "instancia", "plaza", "sede", "y", "no", "nº", "unico", "mixto"}


def court_words(court):
    """Mots qui désignent la ville d'un tribunal : « Sección de lo Mercantil du Tribunal de Instancia de Madrid » -> {'madrid'}."""
    return {w for w in re.findall(r"[a-z]+", fold(court)) if w not in COURT_STOPWORDS and len(w) > 2}


def court_number(court):
    """Numéro de la chambre : « Plaza Nº 16 », « NUMERO 4 DE SEVILLA », « Mercantil 4 de Valencia » ; None s'il n'y en a pas.

    Le BORME écrit parfois « PLAZA Nº 1. 6 » pour 16 (le nombre est coupé) ; le « num. 0 » qui commence tous ses tribunaux est ignoré.
    """
    text = re.sub(r"^\s*num\.?\s*\d+\s+", "", fold(court))
    plaza = re.search(r"plaza\s*n[ºo°]?\.?\s*(\d+(?:\.\s*\d+)?)", text)
    if plaza:
        return int(re.sub(r"\D", "", plaza.group(1)))
    other = re.search(r"(?:numero|n[ºo°])\s*\.?\s*(\d+)|mercantil\s+(\d+)", text)
    return int(other.group(1) or other.group(2)) if other else None


def same_court(lot_court, act_court):
    """True/False si les deux tribunaux se comparent (ville, puis numéro de chambre quand les deux en ont un) ; None sinon."""
    lot_words, act_words = court_words(lot_court), court_words(act_court)
    if not lot_words or not act_words:
        return None
    if not lot_words & act_words:
        return False
    a, b = court_number(lot_court), court_number(act_court)
    return a is None or b is None or a == b


def vinculos(storage, lots=None):
    """Lots rattachés à un acte du BORME : même numéro de procédure ET même tribunal.

    `lots` : les lots à rattacher (défaut : les lots de véhicules ; `es.immobilier` passe ses lots immobiliers).

    Le numéro seul ne suffit pas (chaque tribunal recommence sa numérotation). Le BORME donne le tribunal de chaque acte
    (`juzgado`) : on compare sa ville et son numéro de chambre avec ceux du lot. Sans tribunal d'un côté ou de l'autre, on
    se rabat sur la province du registre de l'acte, citée dans le tribunal ou le lieu du lot (moins précis : la province
    du registre n'est pas forcément celle du tribunal).
    """
    by_procedure = {}
    for acto in storage.rows():
        if acto["procedimiento"]:
            by_procedure.setdefault(normalize_procedure(acto["procedimiento"]), []).append(acto)
    links = []
    for lote in (storage.rows_lotes() if lots is None else lots):
        if not lote["procedimiento"]:
            continue
        where = fold(f"{lote['juzgado'] or ''} {lote['ubicacion'] or ''}")
        companies = {}                                  # une société = une ligne, quel que soit le nombre de ses actes
        for acto in by_procedure.get(normalize_procedure(lote["procedimiento"]), []):
            verdict = same_court(lote["juzgado"], acto["juzgado"])
            if verdict is None:
                provinces = [fold(p) for p in (acto["provincia"] or "").split("/") if p.strip()]
                verdict = any(p in where for p in provinces)
            if verdict:
                companies.setdefault((acto["empresa"], acto["provincia"]), []).append(acto)
        for (empresa, _), actos in companies.items():
            actos.sort(key=lambda a: (a["fecha_publicacion"] or "", a["id"]))          # ordre chronologique, donc stable
            latest = actos[-1]
            admins = {}
            for acto in actos:
                for admin in split_administrators(acto["administradores"]):
                    name, nif = split_admin(admin)
                    admins.setdefault(nif or fold(name), admin)
            links.append({**lote, "empresa": empresa, "administradores": "; ".join(admins.values()) or None,
                          "tipo_acto": latest["tipo"], "url_borme": latest["url_html"], "candidatos": len(companies)})
    return links


# ----------------------------------------------------------------------------------------------------------------------
# Affichage
# ----------------------------------------------------------------------------------------------------------------------

def table(rows, columns=LIQUIDATION_TABLE, width=34, link_key=None):
    """Tableau texte aligné ; `link_key` : le lien est écrit en entier sous chaque ligne (comme côté France)."""
    lines = [[header for _, header in columns]] + [[str(r.get(k) if r.get(k) is not None else "")[:width] for k, _ in columns] for r in rows]
    widths = [max(len(line[i]) for line in lines) for i in range(len(columns))]
    out = ["  ".join(c.ljust(w) for c, w in zip(line, widths)).rstrip() for line in lines]
    out.insert(1, "  ".join("-" * w for w in widths))
    if link_key:
        with_links = out[:2]
        for row, line in zip(rows, out[2:]):
            with_links.append(line)
            if row.get(link_key):
                with_links.append(f"    {row[link_key]}")
        out = with_links
    return "\n".join(out)


def write_csv(rows, path, columns):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    try:
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
    except PermissionError:
        raise SystemExit(f"Impossible d'écrire {path} : le fichier est probablement ouvert dans Excel.")


XLSX_HEADERS = {"first_seen": "vu le", "nb": "nb", "marque": "marque", "municipio": "commune", "tipo": "type de bien",
                "empresa": "société (BORME)", "administradores": "administrateur (BORME)", "url_borme": "texte BORME", "titulo": "lot", "valoracion": "valorisation €",
                "precio_salida": "prix de sortie €", "matricula": "plaque", "ubicacion": "lieu", "juzgado": "tribunal",
                "procedimiento": "procédure", "concursada": "société concursée", "estado": "état", "url": "lien"}
XLSX_WIDTHS = {"titulo": 50, "ubicacion": 30, "juzgado": 40, "concursada": 30, "url": 40}


def write_xlsx(rows, path, columns):
    """Classeur Excel : en-têtes en gras et figés, filtres, colonnes ajustées, liens cliquables, montants numériques."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Véhicules"
    sheet.append([XLSX_HEADERS.get(key, key) for key in columns])
    for row in rows:
        sheet.append([str(row.get(key))[:10] if key == "first_seen" and row.get(key) else row.get(key) for key in columns])
    for number, key in enumerate(columns, start=1):
        letter = get_column_letter(number)
        sheet.column_dimensions[letter].width = XLSX_WIDTHS.get(key, 14)
        for cell in sheet[letter][1:]:
            if key == "url" and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
            elif key in ("valoracion", "precio_salida"):
                cell.number_format = '#,##0 "€"'
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(path)
    except PermissionError:
        raise SystemExit(f"Impossible d'écrire {path} : le fichier est probablement ouvert dans Excel.")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Consulte la base espagnole (BORME + eactivos).")
    parser.add_argument("--db", default=DB_PATH, help="fichier SQLite (défaut : data/es/borme.db)")
    commands = parser.add_subparsers(dest="command", required=True)

    veh = commands.add_parser("vehicles", help="véhicules en vente (lots eactivos)")
    veh.add_argument("--since", help="vus pour la première fois à partir de AAAA-MM-JJ")
    veh.add_argument("--marque", help="filtrer sur une marque (ex : renault)")
    veh.add_argument("--lieu", help="lieu contenant ce texte : ville ou province (ex : madrid, sevilla)")
    veh.add_argument("--min-price", type=float, help="valorisation minimum en euros")
    veh.add_argument("--max-price", type=float, help="valorisation maximum en euros")
    veh.add_argument("--status", choices=list(STATUSES), default="live",
                     help="live (ouverts ou à venir, défaut), open, upcoming, closed ou all")
    veh.add_argument("--sort", choices=SORTS, default="recent", help="recent (défaut) ou price (croissant)")
    veh.add_argument("--limit", type=int, default=30, help="lignes affichées (défaut : 30)")
    veh.add_argument("--csv", help="exporter tous les résultats dans ce fichier CSV")
    veh.add_argument("--xlsx", help="exporter tous les résultats dans ce classeur Excel (.xlsx)")
    veh.add_argument("--no-links", action="store_true", help="ne pas afficher le lien de l'annonce sous chaque véhicule")
    veh.add_argument("--include-unknown", action="store_true", help="garder les véhicules dont la valorisation est inconnue")

    unm = commands.add_parser("unmatched", help="administrateurs concursaux (l'Espagne n'a pas de registre : tous)")
    unm.add_argument("--since", help="à partir de AAAA-MM-JJ")
    unm.add_argument("--limit", type=int, default=25)
    unm.add_argument("--csv", help="exporter les liquidations (pistes) qui ont un administrateur dans ce fichier CSV")

    liq = commands.add_parser("liquidations", help="actes de procédure collective du BORME")
    liq.add_argument("--type", choices=list(TYPES), help="declaration, liquidation ou conclusion")
    liq.add_argument("--lieu", help="province contenant ce texte (sans accents ni majuscules)")
    liq.add_argument("--entreprise", help="nom d'entreprise contenant ce texte")
    liq.add_argument("--since", help="publiés à partir de AAAA-MM-JJ")
    liq.add_argument("--limit", type=int, default=30, help="lignes affichées (défaut : 30)")
    liq.add_argument("--csv", help="exporter tous les résultats dans ce fichier CSV")

    links = commands.add_parser("links", help="lots rattachés à un acte du BORME, donc à un administrateur")
    links.add_argument("--limit", type=int, default=30)
    args = parser.parse_args(argv)

    with Storage(args.db) as storage:
        if args.command == "vehicles":
            found = vehicle_rows(storage, args.status, args.since, args.marque, args.lieu)
            rows = sort_vehicles(filter_price(found, args.min_price, args.max_price, args.include_unknown), args.sort)
            note = ""
            if not args.include_unknown:
                hidden = len(filter_price(found, args.min_price, args.max_price, True)) - len(rows)
                if hidden:
                    note = f"\n{hidden} autre(s) véhicule(s) écarté(s) car une valeur est inconnue (--include-unknown pour les garder)"
            if args.csv or args.xlsx:
                columns = ["codigo", *LOTE_COLUMNS, "marque", "nb", "first_seen"]
                if args.csv:
                    write_csv(rows, args.csv, columns)
                if args.xlsx:
                    write_xlsx(rows, args.xlsx, columns)
                print(f"{len(rows)} véhicule(s) exporté(s) dans {' et '.join(p for p in (args.csv, args.xlsx) if p)}{note}")
            else:
                print(table(rows[:args.limit], VEHICLE_TABLE, width=44, link_key=None if args.no_links else "url_detail")
                      if rows else "Aucun véhicule.")
                total = sum(r["nb"] for r in rows)
                count = f"{len(rows)} véhicule(s)" if total == len(rows) else f"{len(rows)} lot(s), {total} véhicules"
                print(f"\n{count}" + (f", {args.limit} affiché(s)" if len(rows) > args.limit else "") + note)
        elif args.command == "unmatched":
            if args.csv:
                rows = [r for r in select_liquidations(storage, since=args.since) if r["administradores"]]
                write_csv(rows, args.csv, ["id", *COLUMNS])
                print(f"{len(rows)} liquidation(s) avec administrateur exportée(s) dans {args.csv}")
            else:
                ranking = top_administrators(storage, args.since)
                print(table([{"a": a, "n": n} for a, n in ranking[:args.limit]],
                            [("a", "administrateur concursal (à chercher : où annonce-t-il ses biens ?)"), ("n", "procédures")], width=70)
                      if ranking else "Aucun administrateur.")
                print(f"\n{len(ranking)} administrateur(s) distinct(s)")
        elif args.command == "links":
            rows = vinculos(storage)
            print(table(rows[:args.limit], LINK_TABLE, width=40) if rows else "Aucun lot rattaché à un acte du BORME.")
            ambiguous = len({r["codigo"] for r in rows if r["candidatos"] > 1})
            print(f"\n{len({r['codigo'] for r in rows})} lot(s) rattaché(s)"
                  + (f", dont {ambiguous} avec plusieurs sociétés possibles (une ligne par société, colonne « cand. »)" if ambiguous else "")
                  + ". Le BORME collecté ne couvre que ses dates : plus il en couvre, plus il y a de liens.")
        else:
            rows = select_liquidations(storage, args.type, args.lieu, args.since, args.entreprise)
            if args.csv:
                write_csv(rows, args.csv, ["id", *COLUMNS])
                print(f"{len(rows)} ligne(s) écrites dans {args.csv}")
            print(table(rows[:args.limit]) if rows else "Aucun acte.")
            print(f"\n{len(rows)} acte(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
