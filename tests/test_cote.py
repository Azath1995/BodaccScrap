"""`bodacc.cote` : cote saisie à la main, gardée à part des données collectées et affichée par le rapport."""
import pytest

from bodacc import cote, report
from bodacc.storage import Storage


def vehicle(guid, plate):
    return {"guid": guid, "cabinet": "Denis Hazane", "titre": "A vendre : INVENTAIRE CP 60150 MÉLICOCQ", "type": "Véhicules",
            "marque": "Renault", "immatriculation": plate, "kilometrage": 30126, "estimation": "16 000,00 €",
            "lieu": "60150 MÉLICOCQ", "url_detail": f"https://x/{guid}"}


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "base.db"
    with Storage(path) as storage:
        storage.upsert_actifs([vehicle("1", "FX-505-JK"), vehicle("2", "AB-123-CD")], now="2026-10-08T10:00:00+00:00")
    return path


def test_a_cote_is_found_by_plate_and_shown_in_the_export(db, tmp_path):
    cote.main(["fx 505jk", "24408", "--source", "La Centrale", "--db", str(db)])
    out = tmp_path / "v.csv"
    report.main(["--db", str(db), "vehicles", "--include-unknown", "--csv", str(out)])
    text = out.read_text(encoding="utf-8-sig")
    assert "24408" in text and "La Centrale" in text and text.count("24408") == 1


def test_the_cote_survives_an_update_of_the_listing(db):
    cote.main(["FX-505-JK", "24408", "--db", str(db)])
    with Storage(db) as storage:
        storage.upsert_actifs([vehicle("1", "FX-505-JK")], now="2026-10-09T10:00:00+00:00")
        assert storage.cotes()["1"]["valeur"] == 24408


def test_it_can_be_replaced_and_removed(db):
    cote.main(["FX-505-JK", "24408", "--db", str(db)])
    cote.main(["FX-505-JK", "23000", "--db", str(db)])
    with Storage(db) as storage:
        assert storage.cotes()["1"]["valeur"] == 23000
    cote.main(["FX-505-JK", "--supprimer", "--db", str(db)])
    with Storage(db) as storage:
        assert storage.cotes() == {}


def test_cotes_typed_in_the_excel_export_are_imported(db, tmp_path, capsys):
    path = tmp_path / "v.xlsx"
    cote.main(["AB-123-CD", "5000", "--source", "La Centrale", "--db", str(db)])
    report.main(["--db", str(db), "vehicles", "--include-unknown", "--xlsx", str(path)])
    from openpyxl import load_workbook
    book = load_workbook(path)
    sheet = book.active
    headers = [c.value for c in sheet[1]]
    col, link = headers.index("cote €"), headers.index("lien")
    for line in sheet.iter_rows(min_row=2):
        if line[link].value == "https://x/1":
            line[col].value = 22490                         # tapé dans Excel
        else:
            line[col].value = None                          # case vidée : ne supprime rien
    sheet.append([None] * (col + 1) + [None] * (len(headers) - col - 1))
    book.save(path)

    cote.main(["--importer", str(path), "--db", str(db)])
    assert "1 cote(s) enregistrée(s)" in capsys.readouterr().out
    with Storage(db) as storage:
        cotes = storage.cotes()
    assert cotes["1"]["valeur"] == 22490 and cotes["2"]["valeur"] == 5000 and cotes["2"]["source"] == "La Centrale"


def test_an_unreadable_cote_is_explained(db, tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(["lien", "cote € "])
    book.active.append(["https://x/1", "environ 20k"])
    book.save(tmp_path / "v.xlsx")
    with pytest.raises(SystemExit, match="nombre"):
        cote.main(["--importer", str(tmp_path / "v.xlsx"), "--db", str(db)])


def test_an_unknown_or_ambiguous_plate_is_refused(db):
    with pytest.raises(SystemExit, match="Aucun"):
        cote.main(["ZZ-999-ZZ", "1000", "--db", str(db)])
    with Storage(db) as storage:
        storage.upsert_actifs([vehicle("3", "FX-505-JK")])
    with pytest.raises(SystemExit, match="URL"):
        cote.main(["FX-505-JK", "1000", "--db", str(db)])
    cote.main(["https://x/3", "1000", "--db", str(db)])
