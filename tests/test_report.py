import csv

import pytest

from bodacc import report
from bodacc.gemweb_rss import CabinetRegistry
from bodacc.storage import Storage


def vehicle(guid, marque, km=None, annee=None, plaque=None, titre="A vendre : VEHICULE", **extra):
    return {"guid": guid, "cabinet": "Basse", "titre": titre, "type": "Véhicules > Autres",
            "marque": marque, "kilometrage": km, "annee": annee, "immatriculation": plaque,
            "estimation": "1 200,00 €", "lieu": "91000 EVRY", **extra}


@pytest.fixture
def db():
    storage = Storage(":memory:")
    storage.upsert_actifs([
        vehicle("1", "Renault", 174676, 2017, "FX-202-CD"),
        vehicle("2", "Citroën", 90000, 2015),
        {"guid": "3", "cabinet": "Basse", "titre": "Table", "type": "Mobilier"},   # pas un véhicule
    ], now="2026-09-01T10:00:00+00:00")
    storage.upsert_actifs([vehicle("4", "Renault", 5000, 2022)], now="2026-09-10T10:00:00+00:00")
    yield storage
    storage.close()


class TestVehicleRows:
    def test_only_vehicles_newest_first(self, db):
        guids = [r["guid"] for r in report.vehicle_rows(db)]

        assert guids[0] == "4"                      # vu le 10/09, les deux autres le 01/09
        assert set(guids[1:]) == {"1", "2"}         # l'ordre entre égalités n'est pas défini
        assert "3" not in guids                     # une table n'est pas un véhicule

    def test_since_filter(self, db):
        assert [r["guid"] for r in report.vehicle_rows(db, since="2026-09-05")] == ["4"]

    def test_marque_filter_ignores_case_and_accents(self, db):
        assert [r["guid"] for r in report.vehicle_rows(db, marque="citroen")] == ["2"]
        assert {r["guid"] for r in report.vehicle_rows(db, marque="RENAULT")} == {"1", "4"}

    def test_unknown_marque_gives_nothing(self, db):
        assert report.vehicle_rows(db, marque="ferrari") == []


class TestLotsOfSeveralVehicles:
    def lots(self):
        storage = Storage(":memory:")
        storage.upsert_actifs([
            vehicle("10", "Renault", titre="A vendre : Cession de Véhicule 92200", description="deux véhicules CLIO"),
            vehicle("11", "Ford", titre="A vendre : FORD FIESTA 91000", nb_vehicules=1),
        ], now="2026-09-01T10:00:00+00:00")
        return storage

    def test_rows_stored_before_the_column_existed_get_their_count_from_the_text(self):
        counts = {r["guid"]: r["nb_vehicules"] for r in report.vehicle_rows(self.lots())}
        assert counts == {"10": 2, "11": 1}

    def test_a_stored_count_is_kept(self):
        storage = Storage(":memory:")
        storage.upsert_actifs([vehicle("12", "Ford", nb_vehicules=7, description="deux véhicules")], now="2026-09-01T10:00:00+00:00")
        assert report.vehicle_rows(storage)[0]["nb_vehicules"] == 7

    def test_the_summary_counts_vehicles_when_there_are_lots(self, tmp_path, capsys):
        path = tmp_path / "lots.db"
        with Storage(path) as storage:
            storage.upsert_actifs([
                vehicle("10", "Renault", titre="A vendre : Cession de Véhicule 92200", description="deux véhicules CLIO"),
                vehicle("11", "Ford", titre="A vendre : FORD FIESTA 91000")], now="2026-09-01T10:00:00+00:00")
        report.main(["--db", str(path), "vehicles"])
        assert "2 annonce(s), 3 véhicules" in capsys.readouterr().out

    def test_the_summary_is_unchanged_without_lots(self, tmp_path, capsys):
        path = tmp_path / "single.db"
        with Storage(path) as storage:
            storage.upsert_actifs([vehicle("11", "Ford", titre="A vendre : FORD FIESTA 91000")], now="2026-09-01T10:00:00+00:00")
        report.main(["--db", str(path), "vehicles"])
        assert "1 véhicule(s)" in capsys.readouterr().out

    def test_the_table_has_a_count_column(self, db):
        header = report.format_table(report.vehicle_rows(db)).splitlines()[0]
        assert " nb " in header


class TestExcelExport:
    @pytest.fixture
    def path(self, tmp_path):
        target = tmp_path / "base.db"
        with Storage(target) as storage:
            storage.upsert_actifs([
                vehicle("1", "Renault", 174676, 2017, "FX-202-CD", url_detail="https://x/1", description="gasoil"),
                vehicle("2", "Citroën", None, 2015, url_detail="https://x/2"),
            ], now="2026-09-01T10:00:00+00:00")
        return target

    def load(self, xlsx):
        from openpyxl import load_workbook
        return load_workbook(xlsx).active

    def test_the_workbook_has_headers_numbers_and_clickable_links(self, path, tmp_path, capsys):
        target = tmp_path / "vehicules.xlsx"
        report.main(["--db", str(path), "vehicles", "--xlsx", str(target), "--include-unknown"])
        sheet = self.load(target)
        headers = [c.value for c in sheet[1]]
        assert {"vu le", "cabinet", "marque", "km", "lien"} <= set(headers) and sheet.max_row == 3        # en-tête + 2 véhicules
        rows = {r[headers.index("marque")].value: r for r in sheet.iter_rows(min_row=2)}
        renault = rows["Renault"]
        assert renault[headers.index("km")].value == 174676 and isinstance(renault[headers.index("km")].value, int)
        assert renault[headers.index("lien")].hyperlink.target == "https://x/1"
        assert renault[headers.index("vu le")].value == "2026-09-01"
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
        assert "2 véhicule(s) exporté(s)" in capsys.readouterr().out

    def test_filters_apply_to_the_workbook_too(self, path, tmp_path):
        target = tmp_path / "renault.xlsx"
        report.main(["--db", str(path), "vehicles", "--marque", "renault", "--xlsx", str(target)])
        assert self.load(target).max_row == 2

    def test_csv_and_xlsx_can_be_written_together(self, path, tmp_path, capsys):
        csv_path, xlsx_path = tmp_path / "v.csv", tmp_path / "v.xlsx"
        report.main(["--db", str(path), "vehicles", "--csv", str(csv_path), "--xlsx", str(xlsx_path), "--include-unknown"])
        assert csv_path.exists() and xlsx_path.exists() and f"{csv_path} et {xlsx_path}" in capsys.readouterr().out

    def test_a_locked_file_gives_a_clear_message(self, path, tmp_path, monkeypatch):
        def locked(*args, **kwargs):
            raise PermissionError
        monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", locked)
        with pytest.raises(SystemExit, match="ouvert dans Excel"):
            report.main(["--db", str(path), "vehicles", "--xlsx", str(tmp_path / "v.xlsx"), "--include-unknown"])


class TestNumericColumns:
    def test_mileage_sorts_as_a_number_not_as_text(self, db):
        km = [r[0] for r in db.conn.execute(
            "SELECT kilometrage FROM actifs WHERE kilometrage IS NOT NULL ORDER BY kilometrage")]
        assert km == [5000, 90000, 174676]          # en texte, "174676" viendrait avant "5000"
        assert db.conn.execute("SELECT typeof(kilometrage) FROM actifs WHERE guid='1'").fetchone()[0] == "integer"


class TestUnmatchedCabinets:
    def test_counts_by_cabinet_and_uses_the_current_registry(self):
        storage = Storage(":memory:")
        storage.upsert_liquidations([
            {"id": "1", "liquidateur": "la Selarl BLEU SUD Représentée par Me X 8 avenue A 30000 Nîmes"},
            {"id": "2", "liquidateur": "la Selarl BLEU SUD Représentée par Me Y 8 avenue A 30000 Nîmes"},
            {"id": "3", "liquidateur": "Sas les Mandataires Prise en la Personne de Me Z 1 rue A 13300 Salon"},
            {"id": "4", "liquidateur": None},
        ])
        registry = CabinetRegistry({"Les Mandataires": "https://x.fr"})

        assert report.unmatched_cabinets(storage, registry) == [("bleu sud", 2)]

        registry.add("Bleu Sud", "https://bleu.fr")     # le cabinet est ajouté au registre...
        assert report.unmatched_cabinets(storage, registry) == []   # ...il disparaît de la liste
        storage.close()


class TestUnmatchedLiquidations:
    def leads_db(self):
        storage = Storage(":memory:")
        storage.upsert_liquidations([
            {"id": "1", "entreprise": "Boulangerie B", "ville": "Nîmes",
             "liquidateur": "la Selarl BLEU SUD Représentée par Me X 8 avenue A 30000 Nîmes"},
            {"id": "2", "entreprise": "Atelier A", "ville": "Alès",
             "liquidateur": "la Selarl BLEU SUD Représentée par Me Y 8 avenue A 30000 Nîmes"},
            {"id": "3", "entreprise": "Café C", "ville": "Gap",
             "liquidateur": "la Selarl RARE Représentée par Me Z 1 rue B 05000 Gap"},
            {"id": "4", "entreprise": "Déjà connue", "ville": "Salon",
             "liquidateur": "Sas les Mandataires Prise en la Personne de Me W 1 rue C 13300 Salon"},
            {"id": "5", "entreprise": "Sans liquidateur", "liquidateur": None},
        ])
        return storage

    def test_only_unknown_cabinets_frequent_first_with_counts(self):
        rows = report.unmatched_liquidations(
            self.leads_db(), CabinetRegistry({"Les Mandataires": "https://x.fr"}))

        assert [r["entreprise"] for r in rows] == ["Atelier A", "Boulangerie B", "Café C"]
        assert [r["nb_liquidations_cabinet"] for r in rows] == [2, 2, 1]
        assert rows[0]["cabinet_nom"] == "Selarl BLEU SUD"      # nom lisible, pas normalisé

    def test_main_csv_export(self, tmp_path, capsys):
        path, out = tmp_path / "t.db", tmp_path / "pistes.csv"
        with Storage(path) as s:
            s.upsert_liquidations([
                {"id": "1", "entreprise": "Café Été", "ville": "Gap", "url_bodacc": "https://b/1",
                 "liquidateur": "la Selarl ZZZ INCONNU Représentée par Me X 1 rue A 05000 Gap"}])

        report.main(["--db", str(path), "unmatched", "--csv", str(out)])

        assert "1 liquidation(s) hors registre exportée(s)" in capsys.readouterr().out
        with open(out, encoding="utf-8-sig", newline="") as f:
            row = next(csv.DictReader(f))
        assert (row["cabinet_nom"], row["entreprise"], row["nb_liquidations_cabinet"]) == \
               ("Selarl ZZZ INCONNU", "Café Été", "1")
        assert row["url_bodacc"] == "https://b/1"
        assert "_cle" not in row                                 # colonne interne non exportée


class TestOutput:
    def test_table_is_aligned_and_formats_values(self):
        rows = [{"first_seen": "2026-09-19T10:00:00+00:00", "cabinet": "Basse", "marque": "Ford",
                 "annee": None, "kilometrage": 174676, "immatriculation": "FX-101-AB",
                 "carburant": "diesel", "estimation": "1 200,00 €", "lieu": "91000 EVRY"}]

        lines = report.format_table(rows).splitlines()

        assert lines[0].startswith("vu le")
        assert set(lines[1].replace(" ", "")) == {"-"}
        assert lines[2].startswith("2026-09-19  Basse")
        assert "174 676" in lines[2] and "FX-101-AB" in lines[2]

    def test_link_is_written_in_full_under_its_row_and_never_truncated(self):
        url = "https://www.exemple-de-cabinet.fr/anonym/actif/detail/10349186609"        # plus de 30 caractères
        rows = [{"cabinet": "A", "url_detail": url}, {"cabinet": "B"}, {"cabinet": "C", "url_detail": url + "2"}]

        lines = report.format_table(rows, [("cabinet", "cabinet")], link_key="url_detail").splitlines()

        assert lines[2:] == ["A", f"    {url}", "B", "C", f"    {url}2"]     # B n'a pas de lien : pas de ligne vide

    def test_no_link_key_leaves_the_table_unchanged(self):
        rows = [{"cabinet": "A", "url_detail": "https://x"}]
        assert report.format_table(rows, [("cabinet", "cabinet")]) == "cabinet\n-------\nA"

    def test_long_cells_are_truncated(self):
        rows = [{"cabinet": "x" * 100}]
        assert max(len(line) for line in report.format_table(rows, [("cabinet", "cabinet")]).splitlines()) == 30

    def test_csv_keeps_accents_and_all_columns(self, tmp_path):
        rows = [vehicle("1", "Citroën", 90000, 2015, titre="A vendre : CITROËN BERLINGO", description="gasoil")]
        path = report.write_csv(rows, tmp_path / "sub" / "v.csv")

        with open(path, encoding="utf-8-sig", newline="") as f:
            row = next(csv.DictReader(f))

        assert row["marque"] == "Citroën" and row["kilometrage"] == "90000"
        assert row["titre"] == "A vendre : CITROËN BERLINGO" and row["description"] == "gasoil"


class TestMain:
    def test_vehicles_table(self, tmp_path, capsys):
        path = tmp_path / "t.db"
        with Storage(path) as s:
            s.upsert_actifs([vehicle("1", "Renault", 174676, 2017, "FX-202-CD")])

        report.main(["--db", str(path), "vehicles"])

        out = capsys.readouterr().out
        assert "FX-202-CD" in out and "1 véhicule(s)" in out

    def test_limit_is_reported(self, tmp_path, capsys):
        path = tmp_path / "t.db"
        with Storage(path) as s:
            s.upsert_actifs([vehicle(str(i), "Ford") for i in range(5)])

        report.main(["--db", str(path), "vehicles", "--limit", "2"])

        assert "5 véhicule(s), 2 affiché(s)" in capsys.readouterr().out

    def test_csv_export_ignores_the_display_limit(self, tmp_path, capsys):
        path, out_csv = tmp_path / "t.db", tmp_path / "out.csv"
        with Storage(path) as s:
            s.upsert_actifs([vehicle(str(i), "Ford") for i in range(5)])

        report.main(["--db", str(path), "vehicles", "--limit", "2", "--csv", str(out_csv)])

        assert "5 véhicule(s) exporté(s)" in capsys.readouterr().out
        assert len(out_csv.read_text(encoding="utf-8-sig").splitlines()) == 6   # en-tête + 5

    @pytest.mark.parametrize("command", [["vehicles"], ["unmatched"]])
    def test_locked_csv_gives_a_clear_message_not_a_traceback(self, tmp_path, monkeypatch, command):
        def locked(*args, **kwargs):
            raise PermissionError(13, "Permission denied")

        monkeypatch.setattr(report, "write_csv", locked)
        with Storage(tmp_path / "t.db") as s:
            s.upsert_actifs([vehicle("1", "Ford")])

        with pytest.raises(SystemExit) as error:
            report.main(["--db", str(tmp_path / "t.db"), *command, "--csv", str(tmp_path / "x.csv")])

        assert "ouvert dans Excel" in str(error.value)

    def test_unmatched_command(self, tmp_path, capsys):
        path = tmp_path / "t.db"
        with Storage(path) as s:
            s.upsert_liquidations([{"id": "1", "liquidateur": "la Selarl ZZZ INCONNU Représentée par Me X 1 rue A"}])

        report.main(["--db", str(path), "unmatched"])

        assert "zzz inconnu" in capsys.readouterr().out
