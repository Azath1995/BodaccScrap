"""`es.report vehicles` : mêmes options que `bodacc.report vehicles` (marque, lieu, prix, tri, limite, CSV, liens)."""
import csv

import pytest

from es import report
from es.storage import Storage


def lot(code, title, state="abierta", price=None, place=None, court=None, procedure=None, plate=None, seen="2026-09-20", **extra):
    return {"codigo": code, "titulo": title, "estado": state, "valoracion": price, "ubicacion": place, "juzgado": court,
            "procedimiento": procedure, "matricula": plate, "url": f"https://www.eactivos.com/{code}/x", "fecha_sitemap": seen, **extra}


LOTS = [
    lot("AAAAAAA1", "Vehículo Renault Master en Mejorada del Campo (Madrid)", "proxima", 4600, "Mejorada del Campo (Madrid)", "Tribunal de Madrid", "470/2025", "2222CCC"),
    lot("AAAAAAA2", "Vehículo Mitsubishi Outlander en Sevilla", "abierta", 12000, "Sevilla", None, None, None),
    lot("AAAAAAA3", "Furgoneta Dacia Dokker en Madrid_Lote 1", "abierta", 1500, "Madrid", None, "74/2026", "4444FFF"),
    lot("AAAAAAA4", "Vehículo Ford B-MAX en Ávila", "abierta", None, None),
    lot("AAAAAAA5", "Subasta de furgoneta Volkswagen Caddy", "cerrada", 3000),
    lot("AAAAAAA6", "Camión Iveco en La Mojonera (Almería)", "abierta", 1663, None),
]


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "es.db"
    with Storage(path) as s:
        s.upsert_lotes(LOTS, now="2026-09-20T10:00:00+00:00")
        s.upsert_lotes([lot("AAAAAAA7", "Motocicleta Piaggio en Mislata (Valencia)", "proxima", 700, "Mislata (Valencia)")], now="2026-09-10T10:00:00+00:00")
    return path


def run(db, *args):
    report.main(["--db", str(db), "vehicles", *args])


class TestGuessMake:
    @pytest.mark.parametrize("title, make", [
        ("Vehículo Renault Master en Mejorada del Campo (Madrid)", "Renault"), ("Vehículo Mercedes Benz C220D en Calamonte", "Mercedes Benz"),
        ("Vehículo Land Rover en Madrid", "Land Rover"), ("Vehículo Citroën C4 en Bilbao", "Citroen"), ("Motocicleta Piaggio en Mislata", "Piaggio"),
        ("Camión Iveco en La Mojonera", "Iveco"), ("Vehículo Mini Countryman Cooper SD", "Mini"), ("Vivienda en Manzanares", None),
        ("Lote de grúa pluma", None)])
    def test_make_from_the_title(self, title, make):
        assert report.guess_make(title) == make

    def test_a_make_is_a_whole_word(self):
        assert report.guess_make("Vehículo Manuel Kia") == "Kia" and report.guess_make("Terreno en Seattle") is None


class TestCountVehicles:
    @pytest.mark.parametrize("title, count", [
        ("Lote de cinco vehículos en La Romana (Alicante)", 5), ("3 furgonetas en Madrid", 3), ("Dos camiones en Sevilla", 2),
        ("Vehículo Renault Master en Mejorada del Campo (Madrid)", 1),
        ("Furgoneta Dacia Dokker en Madrid_Lote 1", 1),                    # « Lote 1 » : numéro de lot
        ("Lote de vehículos en Madrid", 1),                                # pluriel sans nombre : on n'invente pas
        ("Vivienda en Manzanares", 1)])
    def test_count(self, title, count):
        assert report.count_vehicles(title) == count

    def test_a_number_glued_to_a_word_is_not_a_count(self):
        assert report.count_vehicles("Vehículo 2019 Ford") == 1 and report.count_vehicles("Nave 12coches") == 1


class TestLotsOfSeveralVehicles:
    @pytest.fixture
    def lots_db(self, tmp_path):
        path = tmp_path / "lots.db"
        with Storage(path) as s:
            s.upsert_lotes([lot("BBBBBBB1", "Lote de cinco vehículos en La Romana (Alicante)", "abierta", 9000, "La Romana (Alicante)"),
                            lot("BBBBBBB2", "Vehículo Ford B-MAX en Ávila", "abierta", 2000, "Ávila")], now="2026-09-20T10:00:00+00:00")
        return path

    def test_the_table_shows_the_count_and_the_summary_counts_vehicles_not_lots(self, lots_db, capsys):
        run(lots_db)
        out = capsys.readouterr().out
        assert "nb" in out.splitlines()[0] and "2 lot(s), 6 véhicules" in out

    def test_without_multiple_lots_the_summary_is_unchanged(self, db, capsys):
        run(db)
        assert "6 véhicule(s)" in capsys.readouterr().out

    def test_the_price_stays_the_price_of_the_whole_lot(self, lots_db, capsys):
        run(lots_db, "--max-price", "5000")
        out = capsys.readouterr().out
        assert "Ford B-MAX" in out and "cinco vehículos" not in out       # 9000 € est le prix du lot entier, pas divisé

    def test_csv_has_the_count(self, lots_db, tmp_path):
        target = tmp_path / "v.csv"
        run(lots_db, "--csv", str(target))
        with open(target, encoding="utf-8-sig", newline="") as f:
            assert sorted(int(r["nb"]) for r in csv.DictReader(f)) == [1, 5]


class TestVehicles:
    def test_by_default_only_live_lots_are_listed_newest_seen_first(self, db, capsys):
        run(db)
        out = capsys.readouterr().out
        assert "Renault Master" in out and "Mitsubishi" in out and "Piaggio" in out
        assert "Volkswagen Caddy" not in out                                         # clos : masqué par défaut
        assert out.index("Renault Master") < out.index("Piaggio")                     # vu le 20 septembre avant vu le 10
        assert "6 véhicule(s)" in out

    def test_status_option(self, db, capsys):
        run(db, "--status", "closed")
        assert "Volkswagen Caddy" in capsys.readouterr().out
        run(db, "--status", "all")
        assert "7 véhicule(s)" in capsys.readouterr().out                            # 6 vivants + 1 clos
        run(db, "--status", "upcoming")
        assert "2 véhicule(s)" in capsys.readouterr().out

    def test_marque_filter_ignores_case_and_accents(self, db, capsys):
        run(db, "--marque", "MITSUBISHI")
        out = capsys.readouterr().out
        assert "Outlander" in out and "1 véhicule(s)" in out

    def test_lieu_filter_uses_place_court_and_title(self, db, capsys):
        run(db, "--lieu", "madrid")
        out = capsys.readouterr().out
        assert "2 véhicule(s)" in out and "Renault Master" in out and "Dacia Dokker" in out     # lieu + tribunal, et lieu
        run(db, "--lieu", "MADRID")
        assert "2 véhicule(s)" in capsys.readouterr().out

    def test_price_bounds_are_inclusive_and_unknown_values_are_reported(self, db, capsys):
        run(db, "--max-price", "4600")
        out = capsys.readouterr().out
        assert "Renault Master" in out and "Mitsubishi" not in out and "Dacia" in out
        assert "1 autre(s) véhicule(s) écarté(s) car une valeur est inconnue (--include-unknown pour les garder)" in out     # le Ford sans valorisation

    def test_include_unknown_keeps_lots_without_a_price(self, db, capsys):
        run(db, "--max-price", "4600", "--include-unknown")
        out = capsys.readouterr().out
        assert "Ford B-MAX" in out and "écarté" not in out

    def test_min_price(self, db, capsys):
        run(db, "--min-price", "5000")
        out = capsys.readouterr().out
        assert "Mitsubishi" in out and "Renault" not in out and "1 véhicule(s)" in out

    def test_sort_by_price_puts_unknown_last(self, db, capsys):
        run(db, "--sort", "price")
        out = capsys.readouterr().out
        assert out.index("Piaggio") < out.index("Dacia") < out.index("Iveco") < out.index("Renault") < out.index("Mitsubishi") < out.index("Ford")

    def test_since_uses_the_first_seen_date(self, db, capsys):
        run(db, "--since", "2026-09-15")
        out = capsys.readouterr().out
        assert "Piaggio" not in out and "Renault" in out

    def test_the_link_is_written_under_each_vehicle_unless_disabled(self, db, capsys):
        run(db, "--marque", "renault")
        assert "    https://www.eactivos.com/AAAAAAA1/x" in capsys.readouterr().out
        run(db, "--marque", "renault", "--no-links")
        assert "https://www.eactivos.com" not in capsys.readouterr().out

    def test_limit_says_how_many_are_shown(self, db, capsys):
        run(db, "--limit", "2")
        assert "6 véhicule(s), 2 affiché(s)" in capsys.readouterr().out

    def test_columns_show_make_plate_price_and_procedure(self, db, capsys):
        run(db, "--marque", "renault")
        out = capsys.readouterr().out
        assert "Renault" in out and "2222CCC" in out and "4600" in out and "470/2025" in out and "à venir" in out

    def test_csv_export_has_every_result_and_does_not_truncate(self, db, tmp_path, capsys):
        target = tmp_path / "vehicules.csv"
        run(db, "--csv", str(target), "--limit", "1")
        with open(target, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 6 and {r["marque"] for r in rows} >= {"Renault", "Mitsubishi"}
        assert "6 véhicule(s) exporté(s)" in capsys.readouterr().out

    def test_xlsx_export_has_every_result_with_numbers_and_links(self, db, tmp_path, capsys):
        from openpyxl import load_workbook
        target = tmp_path / "vehiculos.xlsx"
        run(db, "--xlsx", str(target), "--limit", "1")
        sheet = load_workbook(target).active
        headers = [c.value for c in sheet[1]]
        assert sheet.max_row == 7 and {"lot", "valorisation €", "lien", "nb"} <= set(headers) and sheet.freeze_panes == "A2"
        rows = {r[headers.index("lot")].value: r for r in sheet.iter_rows(min_row=2)}
        master = rows["Vehículo Renault Master en Mejorada del Campo (Madrid)"]
        assert master[headers.index("valorisation €")].value == 4600 and master[headers.index("lien")].hyperlink.target.startswith("https://")
        assert "6 véhicule(s) exporté(s)" in capsys.readouterr().out

    def test_empty_database(self, tmp_path, capsys):
        run(tmp_path / "vide.db")
        assert "Aucun véhicule." in capsys.readouterr().out

    def test_unknown_options_of_the_other_country_are_refused(self, db):
        for option in (["--max-km", "50000"], ["--carburant", "diesel"], ["--sort", "km"]):
            with pytest.raises(SystemExit):
                run(db, *option)                            # km, carburant : non extraits des lots espagnols
