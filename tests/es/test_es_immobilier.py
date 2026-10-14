"""`es.immobilier` : lots immobiliers de la région de Torrevieja (eactivos), séparés des véhicules."""
import csv
from datetime import date

import pytest

import test_es_eactivos as ea
from es import eactivos, report
from es.immobilier import report as ireport, update as iupdate, zone
from es.storage import Storage

ROBOTS = ea.ROBOTS
BASE = "https://www.eactivos.com"


def url(code, slug):
    return f"{BASE}/{code}/{slug}"


def sitemap(entries):
    body = "".join(f"<url><loc>{u}</loc><lastmod>{day}</lastmod></url>" for u, day in entries)
    return f'<?xml version="1.0"?><urlset>{body}</urlset>'


def lot_page(title, place, court, procedure, value="85.000", state_text="Finaliza el 30/09/2026 (Europa/Madrid)"):
    body = (f"Portada Liquidaciones Inmuebles {title} {title} {place} Guardar Guardado Compartir Imprimir {state_text} Descripción "
            f"Características Descripción Proced. concursal {procedure} {court} Subasta de un inmueble. Características UBICACIÓN "
            f"{place} VISITAS En espera de información VALORACIÓN {value}€ GASTOS DE GESTIÓN NO INCLUIDOS 5%")
    return ea.page(title, body, 777)


TORREVIEJA = url("AAAAAAA1", "vivienda-duplex-b11-en-torrevieja-alicante")
ORIHUELA = url("AAAAAAA2", "local-comercial-en-orihuela-alicante")
GUARDAMAR = url("AAAAAAA3", "lote-de-trasteros-y-plazas-de-garaje-en-guardamar-del-segura")
VEHICLE = url("AAAAAAA4", "vehiculo-renault-clio-en-torrevieja-alicante")           # un véhicule : pas de l'immobilier
MADRID = url("AAAAAAA5", "vivienda-en-madrid")                                        # hors zone
OLD = url("AAAAAAA6", "vivienda-en-torrevieja-alicante")                              # trop ancien
HOMONYM = url("AAAAAAA7", "vivienda-en-orihuela-del-tremedal-teruel")                 # citée « orihuela » mais en province de Teruel

ENTRIES = [(TORREVIEJA, "2026-09-10"), (ORIHUELA, "2026-09-05"), (GUARDAMAR, "2026-08-20"), (VEHICLE, "2026-09-01"),
           (MADRID, "2026-09-02"), (OLD, "2020-01-01"), (HOMONYM, "2026-09-03")]
PAGES = {
    TORREVIEJA: lot_page("Vivienda dúplex B11 en Torrevieja (Alicante)", "Torrevieja, Alicante", "Juzgado Mercantil 3 de Alicante", "470/2025"),
    ORIHUELA: lot_page("Local comercial en Orihuela (Alicante)", "Orihuela, Alicante", "Juzgado Mercantil 3 de Alicante", "199/2022", "120.000"),
    GUARDAMAR: lot_page("Lote de trasteros y plazas de garaje en Guardamar del Segura (Alicante)", "Guardamar del Segura, Alicante",
                        "Juzgado Mercantil 1 de Alicante", "300/2024", "9.500"),
    HOMONYM: lot_page("Vivienda en Orihuela del Tremedal (Teruel)", "Orihuela del Tremedal, Teruel", "Juzgado de Teruel", "5/2025"),
}


def fake_network(pages=None, entries=None):
    calls = []
    pages = PAGES if pages is None else pages
    entries = ENTRIES if entries is None else entries

    def fetch(address):
        calls.append(address)
        if address.endswith("/robots.txt"):
            return 200, ROBOTS
        if address == eactivos.SITEMAP_URL:
            return 200, sitemap(entries)
        return (200, pages[address]) if address in pages else (404, "")

    fetch.calls = calls
    return fetch


class TestZone:
    @pytest.mark.parametrize("slug, town", [
        ("vivienda-duplex-b11-en-torrevieja-alicante", "Torrevieja"), ("local-en-orihuela-costa-alicante", "Orihuela"),
        ("vivienda-en-cabo-roig", "Orihuela"), ("finca-en-pilar-de-la-horadada", "Pilar de la Horadada"),
        ("vivienda-en-daya-nueva", "Daya Nueva"), ("vivienda-en-madrid", None), ("vivienda-en-torreviejano", None)])   # mot entier
    def test_municipality(self, slug, town):
        assert zone.municipality(slug) == town

    @pytest.mark.parametrize("slug, kind", [
        ("vivienda-duplex-b11-en-torrevieja", "vivienda"), ("local-comercial-en-orihuela", "local"), ("nave-industrial-en-almoradi", "nave/edificio"),
        ("terreno-rustico-en-cox", "terreno/finca"), ("plaza-de-garaje-en-guardamar", "garaje/trastero"),
        ("lote-de-inmuebles-en-torrevieja", "terreno/finca"), ("lote-en-torrevieja", "lote"),
        ("lote-de-maquinaria-en-orihuela", None), ("vehiculo-renault-clio-en-torrevieja", None), ("subasta-de-camion-en-orihuela", None)])
    def test_property_type(self, slug, kind):
        assert zone.property_type(slug) == kind

    def test_a_town_can_be_restricted_by_name(self):
        wanted = zone.keys_for(["Torrevieja", "pilar de la horadada"])
        assert zone.municipality("vivienda-en-torrevieja", wanted) == "Torrevieja" and zone.municipality("vivienda-en-orihuela", wanted) is None

    def test_an_unknown_town_is_refused(self):
        with pytest.raises(ValueError, match="commune inconnue"):
            zone.keys_for(["Madrid"])


class TestSelection:
    def test_only_property_in_the_zone_and_recent(self):
        chosen = iupdate.select(ENTRIES, "2025-01-01", None)
        assert [(u, town, kind) for u, _, town, kind in chosen] == [
            (TORREVIEJA, "Torrevieja", "vivienda"), (ORIHUELA, "Orihuela", "local"), (HOMONYM, "Orihuela", "vivienda"),
            (GUARDAMAR, "Guardamar del Segura", "garaje/trastero")]        # trié du plus récent au plus ancien ; ni véhicule, ni Madrid, ni 2020

    def test_known_closed_lots_are_not_read_again(self):
        chosen = iupdate.select(ENTRIES, "2025-01-01", None, known_closed={"AAAAAAA1"})
        assert TORREVIEJA not in [u for u, *_ in chosen]


class TestProvince:
    @pytest.mark.parametrize("title, place, expected", [
        ("Vivienda en Torrevieja (Alicante)", "Torrevieja, Alicante", True),
        ("Vivienda en Torrevieja (Alicante)", None, True),                                     # fiche close : le titre le dit
        ("Vivienda 5ª G en Guardamar de la Safor (Valencia)", None, False),                    # homonyme en province de Valence
        ("Subasta de local comercial en Villamartín (Cádiz)", None, False),
        ("Vivienda en Orihuela del Tremedal (Teruel)", "Orihuela del Tremedal, Teruel", False),
        ("Finca urbana en Torrevieja", None, True)])                                            # aucune province citée : gardé
    def test_in_province(self, title, place, expected):
        assert iupdate.in_province({"titulo": title, "ubicacion": place}) is expected


class TestUpdate:
    def run(self, tmp_path, *args, fetch=None):
        path = tmp_path / "es.db"
        fetch = fetch or fake_network()
        iupdate.main(["--db", str(path), "--pause", "0", *args], fetch=fetch, sleep=lambda s: None, today=date(2026, 9, 21))
        return path, fetch

    def test_stores_the_zone_property_and_only_that(self, tmp_path, capsys):
        path, fetch = self.run(tmp_path)
        out = capsys.readouterr().out
        with Storage(path) as s:
            rows = {r["codigo"]: r for r in s.rows_inmuebles()}
            assert s.count("lotes") == 0                                        # jamais mélangé avec les véhicules
        assert set(rows) == {"AAAAAAA1", "AAAAAAA2", "AAAAAAA3"}
        assert rows["AAAAAAA1"]["municipio"] == "Torrevieja" and rows["AAAAAAA1"]["tipo"] == "vivienda"
        assert rows["AAAAAAA1"]["valoracion"] == 85000 and rows["AAAAAAA1"]["procedimiento"] == "470/2025"
        assert "1 écarté(s) (hors province d'Alicante)" in out and "3 lus, 3 nouveaux" in out
        assert VEHICLE not in fetch.calls and MADRID not in fetch.calls and OLD not in fetch.calls     # jamais demandés

    def test_nothing_is_requested_that_robots_forbids(self, tmp_path):
        forbidden = url("AAAAAAA8", "vivienda-en-torrevieja-alicante/mejor-puja")
        fetch = fake_network(entries=[(forbidden, "2026-09-10")], pages={})
        self.run(tmp_path, fetch=fetch)
        assert forbidden not in fetch.calls

    def test_second_run_adds_nothing_and_keeps_history(self, tmp_path, capsys):
        path, _ = self.run(tmp_path)
        capsys.readouterr()
        self.run(tmp_path)
        assert "0 nouveaux" in capsys.readouterr().out
        with Storage(path) as s:
            assert s.count("inmuebles") == 3

    def test_since_and_municipios_options(self, tmp_path):
        path, _ = self.run(tmp_path, "--since", "2026-09-06", "--municipios", "torrevieja")
        with Storage(path) as s:
            assert [r["codigo"] for r in s.rows_inmuebles()] == ["AAAAAAA1"]

    @pytest.mark.parametrize("args", [["--since", "21/09/2026"], ["--pause", "-1"], ["--municipios", "madrid"]])
    def test_invalid_options_are_refused(self, tmp_path, args):
        with pytest.raises(SystemExit):
            iupdate.main(["--db", str(tmp_path / "x.db"), *args], fetch=fake_network(), sleep=lambda s: None)


def borme_act(**extra):
    base = {"id": "a1", "fecha_publicacion": "2026-07-14", "provincia": "ALICANTE/ALACANT", "borme_id": "BORME-A-2026-133-03",
            "empresa_num": "1", "empresa": "PARQUE DE LAS NACIONES DE TORREVIEJA SL", "procedimiento": "470/2025",
            "tipo": "Auto de apertura de la fase de liquidación", "administradores": "PKF ATTEST SERVICIOS EMPRESARIALES SL (B12345678)",
            "juzgado": "num. 0 JUZGADO DE LO MERCANTIL NUMERO 3 DE ALICANTE", "url_html": "https://h/1"}
    return {**base, **extra}


class TestReport:
    @pytest.fixture
    def db(self, tmp_path):
        path, _ = TestUpdate().run(tmp_path)
        return path

    def run(self, db, *args):
        """`links ...` tel quel ; tout le reste est une option de la sous-commande `lots`."""
        command = [] if args and args[0] == "links" else ["lots"]
        ireport.main(["--db", str(db), *command, *args])

    def test_lists_live_lots_by_default_with_town_and_type(self, db, capsys):
        self.run(db)
        out = capsys.readouterr().out
        assert "Torrevieja" in out and "vivienda" in out and "Orihuela" in out and "3 lot(s)" in out
        assert "Renault Clio" not in out                                         # un véhicule n'est jamais dans cette liste

    def test_filters(self, db, capsys):
        self.run(db, "--municipio", "torrevieja")
        assert "1 lot(s)" in capsys.readouterr().out
        self.run(db, "--tipo", "local")
        out = capsys.readouterr().out
        assert "Local comercial" in out and "1 lot(s)" in out
        self.run(db, "--max-price", "10000")
        assert "Lote de trasteros" in capsys.readouterr().out
        self.run(db, "--min-price", "100000")
        assert "Local comercial" in capsys.readouterr().out

    def test_closed_lots_are_hidden_unless_asked(self, tmp_path, capsys):
        closed = url("AAAAAAA9", "vivienda-en-torrevieja-alicante")
        page = ea.page("Vivienda en Torrevieja (Alicante)", "Portada Liquidaciones Inmuebles Vaya, la subasta se ha cerrado. VALORACIÓN 60.000€", 1)
        path, _ = TestUpdate().run(tmp_path, fetch=fake_network({closed: page}, [(closed, "2026-09-10")]))
        self.run(path)
        assert "Aucun lot immobilier." in capsys.readouterr().out
        self.run(path, "--status", "closed")
        assert "1 lot(s)" in capsys.readouterr().out

    def test_the_link_is_written_under_each_lot_unless_disabled(self, db, capsys):
        self.run(db, "--municipio", "torrevieja")
        assert f"    {TORREVIEJA}" in capsys.readouterr().out
        self.run(db, "--municipio", "torrevieja", "--no-links")
        assert BASE not in capsys.readouterr().out

    def test_unknown_price_is_reported_and_can_be_kept(self, tmp_path, capsys):
        no_price = url("AAAAAAB1", "vivienda-en-torrevieja-alicante")
        page = ea.page("Vivienda en Torrevieja (Alicante)", "Portada Liquidaciones Inmuebles UBICACIÓN Torrevieja, Alicante VISITAS", 2)
        path, _ = TestUpdate().run(tmp_path, fetch=fake_network({no_price: page}, [(no_price, "2026-09-10")]))
        self.run(path, "--max-price", "100000")
        assert "1 autre(s) lot(s) écarté(s)" in capsys.readouterr().out
        self.run(path, "--max-price", "100000", "--include-unknown")
        assert "1 lot(s)" in capsys.readouterr().out and "écarté" not in capsys.readouterr().out

    def test_csv_and_xlsx_have_every_lot(self, db, tmp_path, capsys):
        from openpyxl import load_workbook
        csv_path, xlsx_path = tmp_path / "i.csv", tmp_path / "i.xlsx"
        self.run(db, "--csv", str(csv_path), "--xlsx", str(xlsx_path), "--limit", "1")
        with open(csv_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 3 and {r["municipio"] for r in rows} == {"Torrevieja", "Orihuela", "Guardamar del Segura"}
        sheet = load_workbook(xlsx_path).active
        headers = [c.value for c in sheet[1]]
        assert sheet.max_row == 4 and "commune" in headers and "type de bien" in headers

    def test_links_join_lots_to_the_borme_company_and_administrator(self, db, capsys):
        with Storage(db) as s:
            s.upsert([borme_act()])
        self.run(db, "links")
        out = capsys.readouterr().out
        assert "PARQUE DE LAS NACIONES DE TORREVIEJA SL" in out and "PKF ATTEST SERVICIOS EMPRESARIALES SL" in out    # colonne tronquée : le CSV a le NIF
        assert "1 lot(s) rattaché(s), sur 3 lot(s) immobilier(s)." in out

    def test_links_ignore_the_vehicle_lots(self, db, capsys):
        with Storage(db) as s:
            s.upsert_lotes([{"codigo": "VVVVVVV1", "titulo": "Vehículo Ford", "estado": "abierta", "procedimiento": "470/2025",
                             "juzgado": "Juzgado Mercantil 3 de Alicante", "ubicacion": "Torrevieja (Alicante)", "url": "https://x/v"}])
            s.upsert([borme_act()])
        assert {r["codigo"] for r in report.vinculos(Storage(db))} == {"VVVVVVV1"}                     # le rapport des véhicules
        self.run(db, "links")
        assert "Vehículo Ford" not in capsys.readouterr().out                                          # pas celui de l'immobilier

    def test_links_can_be_exported(self, db, tmp_path):
        with Storage(db) as s:
            s.upsert([borme_act()])
        target = tmp_path / "links.csv"
        self.run(db, "links", "--csv", str(target))
        with open(target, encoding="utf-8-sig", newline="") as f:
            (row,) = list(csv.DictReader(f))
        assert row["empresa"].startswith("PARQUE") and row["url_borme"] == "https://h/1" and row["municipio"] == "Torrevieja"
        assert row["administradores"] == "PKF ATTEST SERVICIOS EMPRESARIALES SL (B12345678)"

    def test_an_unknown_town_is_refused(self, db):
        with pytest.raises(SystemExit):
            self.run(db, "--municipio", "madrid")
